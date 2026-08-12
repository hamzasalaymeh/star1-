#!/usr/bin/env python3
"""
تشغيل كامل لقرار تسجيل واحد:
1. تسجيل دخول + الوصول لصفحة البحث
2. إدخال رقم القرار والبحث
3. جمع كل روابط الاستمارات من كل صفحات النتائج (بالتنقل بين الصفحات)
4. فتح كل استمارة مباشرة عبر رابطها وتصدير الجدول - بدون الرجوع للقائمة أو إعادة البحث
5. بناء فهرس Excel فيه اسم كل موقع مربوط برابط نسبي لملف الـ PDF الخاص فيه
"""

import asyncio
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from rich import box
from rich.console import Console
from rich.panel import Panel
from rich.table import Table

from table_pdf_exporter_advanced import AdvancedTablePDFExporter

console = Console()

BANNER = r"""[bold green]
   ██████╗████████╗ █████╗ ██████╗  ██╗
  ██╔════╝╚══██╔══╝██╔══██╗██╔══██╗███║
  ╚█████╗    ██║   ███████║██████╔╝╚██║
   ╚═══██╗   ██║   ██╔══██║██╔══██╗ ██║
  ██████╔╝   ██║   ██║  ██║██║  ██║ ██║
  ╚═════╝    ╚═╝   ╚═╝  ╚═╝╚═╝  ╚═╝ ╚═╝
[/bold green][bold cyan]     HERITAGE SURVEY :: CLASSIFICATION TABLE EXPORTER[/bold cyan]
[dim]     تصدير جداول معايير التصنيف إلى PDF + فهرس Excel[/dim]
"""


def print_banner():
    console.print(BANNER)
    console.rule(style="green")


def build_excel_index(output_dir: str, decision_number: str, results: list) -> str:
    """
    ينشئ ملف Excel فيه اسم كل موقع مربوط بهايبرلينك لملف الـ PDF الخاص فيه.
    الرابط نسبي (relative) وليس مسار مطلق، فهو يشتغل من أي جهاز أو مسار
    طالما مجلد الـ exports كامل (بمجلداته الفرعية) بينتقل معاً.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'قرار {decision_number}'[:31]
    ws.sheet_view.rightToLeft = True

    headers = ['#', 'اسم الموقع', 'رابط PDF', 'الحالة']
    header_fill = PatternFill(start_color='1B5E20', end_color='1B5E20', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=12)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    success_fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
    fail_fill = PatternFill(start_color='FFEBEE', end_color='FFEBEE', fill_type='solid')

    row_num = 2
    for i, r in enumerate(results, 1):
        is_success = r['status'] == 'success'
        fill = success_fill if is_success else fail_fill

        ws.cell(row=row_num, column=1, value=i).fill = fill
        name_cell = ws.cell(row=row_num, column=2, value=r.get('siteName') or r.get('name'))
        name_cell.fill = fill
        name_cell.alignment = Alignment(horizontal='right')

        link_cell = ws.cell(row=row_num, column=3)
        link_cell.fill = fill
        if is_success and r.get('relPdfPath'):
            link_cell.value = '📄 فتح الملف'
            link_cell.hyperlink = r['relPdfPath'].replace('\\', '/')
            link_cell.font = Font(color='1565C0', underline='single')
            link_cell.alignment = Alignment(horizontal='center')

        status_cell = ws.cell(row=row_num, column=4,
                               value='✅ نجح' if is_success else f'❌ {r.get("error", "فشل")}')
        status_cell.fill = fill
        status_cell.alignment = Alignment(horizontal='center')

        row_num += 1

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 55
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 30
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:D{row_num - 1}'

    excel_path = os.path.join(output_dir, f'فهرس_قرار_{decision_number}.xlsx')
    wb.save(excel_path)
    return excel_path


async def main():
    print_banner()

    if len(sys.argv) < 2:
        console.print('[bold red]استخدام:[/bold red] python run_decision_export.py <رقم_القرار> [عدد الاستمارات المطلوب] [username] [password]')
        console.print('  [dim]مثال: python run_decision_export.py 21          → يصدر كل الاستمارات[/dim]')
        console.print('  [dim]مثال: python run_decision_export.py 21 10       → يصدر أول 10 استمارات فقط[/dim]')
        sys.exit(1)

    decision_number = sys.argv[1]
    forms_limit = int(sys.argv[2]) if len(sys.argv) > 2 and sys.argv[2].isdigit() else None

    username = sys.argv[3] if len(sys.argv) > 3 else console.input('[bold cyan]👤 اسم المستخدم:[/bold cyan] ')
    password = sys.argv[4] if len(sys.argv) > 4 else console.input('[bold cyan]🔑 كلمة المرور:[/bold cyan] ', password=True)

    exporter = AdvancedTablePDFExporter(
        username,
        password,
        options={
            'headless': True,  # شغّل headless=False فقط لو بدك تراقب المتصفح مباشرة
            'outputDir': f'./exports/قرار_{decision_number}',
        }
    )

    results = []

    try:
        await exporter.init()
        await exporter.login()
        await exporter.goto_search_page()

        # إدخال رقم القرار والبحث
        await exporter.select_decision_number(decision_number)
        await exporter.click_search_button()
        await asyncio.sleep(2)

        # جمع روابط الاستمارات - لو فيه حد محدد، يوقف التنقل بمجرد جمع عدد كافي
        form_links = await exporter.collect_all_form_links(min_links=forms_limit)

        if forms_limit:
            form_links = form_links[:forms_limit]
            console.print(f'[cyan]ℹ️ تم تحديد المعالجة لأول {forms_limit} استمارة فقط[/cyan]')

        if not form_links:
            console.print('[bold red]❌ لم يتم العثور على أي استمارات لهذا القرار[/bold red]')
            await exporter.close()
            return

        total = len(form_links)
        console.rule(f'[bold cyan]جارٍ تصدير {total} استمارة[/bold cyan]', style="cyan")

        for i, link in enumerate(form_links, 1):
            try:
                # رابط تبويب "معايير التصنيف" مباشرة: DataEdit/{id}/10
                site_id = link.rstrip('/').split('/')[-1]
                classification_link = link.rstrip('/') + '/10'
                console.print(f'\n[bold]\\[{i}/{total}][/bold] [dim]{classification_link}[/dim]')

                # فتح تبويب معايير التصنيف مباشرة - بدون فتح الاستمارة الأساسية أو الرجوع لقائمة البحث
                await exporter.page.goto(classification_link, wait_until=exporter.config['waitForNavigation'])

                # extract_red_box_table() ينتظر بنفسه تحميل الجدول وبطاقة التصنيف
                table_info = await exporter.extract_red_box_table()
                site_name = await exporter.get_site_name(site_id=site_id)

                # مجلد خاص بكل موقع يحتوي على الـ PDF والصورة بنفس اسم الموقع
                site_folder = os.path.join(exporter.config['outputDir'], site_name)
                os.makedirs(site_folder, exist_ok=True)

                rel_pdf_path = os.path.join(site_name, f'{site_name}.pdf')
                rel_png_path = os.path.join(site_name, f'{site_name}.png')

                pdf_path = await exporter.export_table_as_clean_pdf(rel_pdf_path, table_info)
                png_path = await exporter.capture_table_screenshot(rel_png_path, table_info)

                results.append({
                    'name': site_name,
                    'siteName': site_name,
                    'status': 'success',
                    'pdfPath': pdf_path,
                    'pngPath': png_path,
                    'relPdfPath': rel_pdf_path,
                })
                console.print(f'  [bold green]✔ نجح:[/bold green] {site_name}')

            except Exception as e:
                console.print(f'  [bold red]✘ فشل الاستمارة {i}: {e}[/bold red]')
                results.append({'name': f'استمارة_{i}', 'status': 'failed', 'error': str(e)})

        await exporter.close()

        # جدول ملخص ملون
        summary = Table(title='النتيجة النهائية', box=box.ROUNDED, show_lines=False,
                         title_style='bold cyan', header_style='bold white on dark_green')
        summary.add_column('اسم الموقع', style='white')
        summary.add_column('الحالة', justify='center')
        for r in results:
            status_text = '[bold green]✅ نجح[/bold green]' if r['status'] == 'success' else '[bold red]❌ فشل[/bold red]'
            summary.add_row(r.get('siteName') or r['name'], status_text)
        console.print(summary)

        success_count = sum(1 for r in results if r['status'] == 'success')
        panel_style = 'green' if success_count == len(results) else 'yellow'
        console.print(Panel(f'[bold]{success_count}/{len(results)}[/bold] استمارة تم تصديرها بنجاح',
                             style=panel_style, box=box.DOUBLE))

        # فهرس Excel بروابط نسبية تشتغل من أي جهاز
        excel_path = build_excel_index(exporter.config['outputDir'], decision_number, results)
        console.print(f'[bold cyan]📊 تم إنشاء فهرس Excel:[/bold cyan] {excel_path}')

    except Exception as e:
        await exporter.close()
        console.print(f'[bold red]❌ خطأ عام: {e}[/bold red]')


if __name__ == '__main__':
    asyncio.run(main())
